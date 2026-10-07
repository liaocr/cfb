#!/usr/bin/env python3
"""Build an optional Excel review surface; proposals are not gold until a single AI review is recorded."""
import argparse
import hashlib
import json
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_QUEUE = ROOT / "transfer/models/micro-generator-review-queue.jsonl"
DEFAULT_PROPOSALS = ROOT / "transfer/models/micro-generator-fact-draft-suggestions.jsonl"
DEFAULT_OUTPUT = ROOT / "transfer/models/micro-generator-ai-review-workbook.xlsx"


def read_jsonl(path):
    rows = []
    with Path(path).open(encoding="utf-8") as handle:
        for line_no, line in enumerate(handle, 1):
            if not line.strip():
                continue
            try:
                rows.append(json.loads(line))
            except Exception as exc:
                raise ValueError(f"{path}:{line_no}: invalid JSON: {exc}") from exc
    return rows


def build_workbook(queue, proposals, output_path):
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise SystemExit("optional review workbook needs openpyxl>=3.1") from exc

    if len(queue) != len(proposals) or {r["caseId"] for r in queue} != {r["caseId"] for r in proposals}:
        raise ValueError("queue and fact-proposal cases do not match")
    for row in [*queue, *proposals]:
        for key in ("raw", "ctx", "referenceDraft"):
            if len(row.get(key, "")) > 32767:
                raise ValueError(f"{row['caseId']}.{key} exceeds the Excel cell text limit")

    wb = Workbook()
    start = wb.active
    start.title = "Start Here"
    cases_ws = wb.create_sheet("Cases")
    facts_ws = wb.create_sheet("Fact Proposals")
    conflicts_ws = wb.create_sheet("Target Conflicts")
    splits_ws = wb.create_sheet("Family Split")

    navy = "17365D"
    blue = "D9EAF7"
    yellow = "FFF2CC"
    red = "FCE4D6"
    green = "E2F0D9"
    thin_gray = Side(style="thin", color="D9E1F2")
    header_fill = PatternFill("solid", fgColor=navy)
    header_font = Font(color="FFFFFF", bold=True)
    wrap_top = Alignment(vertical="top", wrap_text=True)

    def setup_table(ws, headers, widths, freeze="A2"):
        ws.append(headers)
        ws.freeze_panes = freeze
        ws.auto_filter.ref = f"A1:{chr(64 + len(headers))}1"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        ws.row_dimensions[1].height = 34
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

    start.append(["Micro-generator AI single-review workbook"])
    start["A1"].font = Font(size=16, bold=True, color=navy)
    instructions = [
        ["Status", "Review aid only. The workbook itself does not certify any row or make it training-eligible."],
        ["Source facts", "The 70 highlighted suggestions were copied from explicit task/context preamble spans only. They are partial proposals, not a complete reading of raw/context."],
        ["Required review", "Read the complete raw and ctx in the Cases sheet. Correct/add/remove facts, verify exact source spans, importance, conditions, polarity, modality, and must-preserve decisions."],
        ["Single-review policy", "One authorized AI reviewer is sufficient under this project policy. Record reviewerId=arena-agent-mode, reviewerType=ai, reviewerCount=1, humanReviewerCount=0, independentSecondReview=false, date, scope, and limitations in JSON. Do not imply a second or human review."],
        ["Target conflict", "One identical raw+ctx input has two different reference drafts. Compare both in Target Conflicts and choose/merge/exclude only after reviewing source evidence."],
        ["Family split", "Current queued families all have existing-dev provenance. The suggested dev placement is only a conservative proposal; finalSplit remains blank. No train family or new blind family is registered."],
        ["No blind claims", "Do not put existing holdout families into blind. A genuinely new family with private raw+ctx samples is still required."],
        ["No training", "Do not export or train from this workbook. Return completed decisions as annotation JSONL, validate with the schemas, then run the fail-closed exporter."],
        ["Span convention", "Offsets are JavaScript UTF-16 code units, half-open [start,end); verify quote exactly. The proposals were generated and range-checked with JavaScript string offsets."],
        ["Guide", "tools/micro-generator/REVIEW-GUIDE.md"],
        ["Queue", "transfer/models/micro-generator-review-queue.jsonl"],
        ["Draft proposals", "transfer/models/micro-generator-fact-draft-suggestions.jsonl"],
    ]
    for row in instructions:
        start.append(row)
    start.column_dimensions["A"].width = 24
    start.column_dimensions["B"].width = 120
    start.freeze_panes = "A2"
    for row in start.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color=navy)
        row[0].alignment = wrap_top
        row[1].alignment = wrap_top
        for cell in row:
            cell.border = Border(bottom=thin_gray)
        start.row_dimensions[row[0].row].height = 42
    start["B2"].fill = PatternFill("solid", fgColor=red)
    start["B6"].fill = PatternFill("solid", fgColor=yellow)
    start["B8"].fill = PatternFill("solid", fgColor=red)

    case_headers = ["caseId", "family", "sourceKind", "provenanceSplit", "proposedSplit", "finalSplit",
        "sourceCoverage", "annotationStatus", "AI reviewer ID", "AI review provenance", "referenceReview status",
        "unsupportedClaimsReviewed", "unsupportedClaimsCount", "targetConflict?", "rawSha256", "ctxSha256",
        "referenceSha256", "raw", "ctx", "referenceDraft", "reviewer notes"]
    case_widths = [32, 22, 32, 18, 18, 20, 18, 18, 20, 20, 22, 24, 24, 18, 66, 66, 66, 70, 70, 70, 60]
    setup_table(cases_ws, case_headers, case_widths)
    pair_to_cases = defaultdict(list)
    for row in queue:
        raw_hash = row["sourceHashes"]["rawSha256"]
        ctx_hash = row["sourceHashes"]["ctxSha256"]
        pair_to_cases[(raw_hash, ctx_hash)].append(row)
    conflict_case_ids = {r["caseId"] for rows in pair_to_cases.values() if len({x["sourceHashes"]["referenceSha256"] for x in rows}) > 1 for r in rows}
    for row in queue:
        cases_ws.append([row["caseId"], row["family"], row["sourceKind"], row["provenanceSplit"],
            "dev (proposal: preserve existing-dev provenance)", None, row["sourceCoverage"], row["annotationStatus"],
            None, None, row["referenceReview"]["status"], row["referenceReview"]["unsupportedClaimsReviewed"],
            row["referenceReview"]["unsupportedClaimsCount"], "YES—unresolved" if row["caseId"] in conflict_case_ids else "",
            row["sourceHashes"]["rawSha256"], row["sourceHashes"]["ctxSha256"], row["sourceHashes"]["referenceSha256"],
            row["raw"], row["ctx"], row["referenceDraft"], None])
        idx = cases_ws.max_row
        for cell in cases_ws[idx]:
            cell.alignment = wrap_top
            cell.border = Border(bottom=thin_gray)
        cases_ws.cell(idx, 5).fill = PatternFill("solid", fgColor=yellow)
        cases_ws.cell(idx, 6).fill = PatternFill("solid", fgColor=yellow)
        cases_ws.cell(idx, 21).fill = PatternFill("solid", fgColor=yellow)
        cases_ws.row_dimensions[idx].height = 105
    case_split_dv = DataValidation(type="list", formula1='"train,dev,blind"', allow_blank=True)
    coverage_dv = DataValidation(type="list", formula1='"unreviewed,partial,complete"', allow_blank=True)
    status_dv = DataValidation(type="list", formula1='"unreviewed,draft,reviewed,adjudicated"', allow_blank=True)
    cases_ws.add_data_validation(case_split_dv); case_split_dv.add(f"F2:F{cases_ws.max_row}")
    cases_ws.add_data_validation(coverage_dv); coverage_dv.add(f"G2:G{cases_ws.max_row}")
    cases_ws.add_data_validation(status_dv); status_dv.add(f"H2:H{cases_ws.max_row}")

    fact_headers = ["caseId", "family", "factId", "candidate proposition", "source field", "start UTF-16", "end UTF-16", "exact quote",
        "suggested importance", "suggested mustPreserve", "suggested polarity", "suggested modality", "entities", "tags",
        "AI reviewer importance", "AI reviewer mustPreserve", "AI reviewer reference status", "span verified?", "AI reviewer notes"]
    fact_widths = [32, 22, 42, 64, 14, 14, 14, 80, 20, 22, 20, 20, 36, 30, 22, 24, 26, 18, 60]
    setup_table(facts_ws, fact_headers, fact_widths)
    proposal_by_id = {row["caseId"]: row for row in proposals}
    for case in queue:
        proposal = proposal_by_id[case["caseId"]]
        for fact in proposal["facts"]:
            facts_ws.append([case["caseId"], case["family"], fact["factId"], fact["proposition"], fact["sourceSpan"]["field"],
                fact["sourceSpan"]["start"], fact["sourceSpan"]["end"], fact["sourceSpan"]["quote"], fact["importance"],
                fact["mustPreserve"], fact["polarity"], fact["modality"], ", ".join(fact["entities"]), ", ".join(fact.get("tags", [])),
                None, None, None, None, None])
            idx = facts_ws.max_row
            for cell in facts_ws[idx]:
                cell.alignment = wrap_top
                cell.border = Border(bottom=thin_gray)
            for col in (15, 16, 17, 18, 19):
                facts_ws.cell(idx, col).fill = PatternFill("solid", fgColor=yellow)
            facts_ws.row_dimensions[idx].height = 64
    importance_dv = DataValidation(type="list", formula1='"critical,high,medium,low"', allow_blank=True)
    boolean_dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    fact_status_dv = DataValidation(type="list", formula1='"preserved,omitted,contradicted,uncertain"', allow_blank=True)
    facts_ws.add_data_validation(importance_dv)
    for col in (15,): importance_dv.add(f"{chr(64+col)}2:{chr(64+col)}{facts_ws.max_row}")
    facts_ws.add_data_validation(boolean_dv)
    for col in (16, 18): boolean_dv.add(f"{chr(64+col)}2:{chr(64+col)}{facts_ws.max_row}")
    facts_ws.add_data_validation(fact_status_dv)
    for col in (17,): fact_status_dv.add(f"{chr(64+col)}2:{chr(64+col)}{facts_ws.max_row}")
    facts_ws["I1"].comment = Comment("Assistant proposal only; reviewer must decide importance.", "Arena Agent")
    facts_ws["J1"].comment = Comment("Assistant proposal only; reviewer must decide mustPreserve.", "Arena Agent")

    conflicts_ws.append(["Unresolved identical-input / different-target groups"])
    conflicts_ws["A1"].font = Font(size=14, bold=True, color=navy)
    conflict_headers = ["inputSha256", "family", "caseIds", "targetSha256 values", "candidate drafts (verbatim)", "review outcome", "canonical target hash", "AI reviewer ID", "notes"]
    conflicts_ws.append(conflict_headers)
    for cell in conflicts_ws[2]:
        cell.fill = header_fill; cell.font = header_font; cell.alignment = wrap_top
    conflicts_ws.freeze_panes = "A3"
    conflicts_ws.auto_filter.ref = f"A2:I{2+len(conflict_case_ids)}"
    for col, width in enumerate([66, 22, 72, 130, 100, 28, 66, 34, 60], 1):
        conflicts_ws.column_dimensions[chr(64+col)].width = width
    for (raw_hash, ctx_hash), rows in pair_to_cases.items():
        targets = {r["sourceHashes"]["referenceSha256"] for r in rows}
        if len(targets) <= 1:
            continue
        first = rows[0]
        drafts = "\n\n".join(f"[{r['caseId']}]\n{r['referenceDraft']}" for r in rows)
        conflicts_ws.append([hashlib.sha256((raw_hash + ":" + ctx_hash).encode()).hexdigest(), first["family"],
            "\n".join(r["caseId"] for r in rows), "\n".join(sorted(targets)), drafts, None, None, None, None])
        idx = conflicts_ws.max_row
        for cell in conflicts_ws[idx]:
            cell.alignment = wrap_top; cell.border = Border(bottom=thin_gray)
        conflicts_ws.cell(idx, 6).fill = PatternFill("solid", fgColor=yellow)
        conflicts_ws.cell(idx, 9).fill = PatternFill("solid", fgColor=yellow)
        conflicts_ws.row_dimensions[idx].height = 190
    conflict_dv = DataValidation(type="list", formula1='"choose-A,choose-B,merge-after-review,exclude-both,pending"', allow_blank=True)
    conflicts_ws.add_data_validation(conflict_dv); conflict_dv.add(f"F3:F{conflicts_ws.max_row}")

    split_headers = ["family", "source provenance", "proposed split", "finalSplit", "new blind eligible?", "review status", "notes"]
    setup_table(splits_ws, split_headers, [24, 24, 22, 22, 24, 22, 90])
    known = {row["family"] for row in queue}
    provenance = {row["family"]: row["provenanceSplit"] for row in queue}
    for family in sorted(known):
        splits_ws.append([family, provenance[family], "dev (proposal)", None, "NO—already known", "pending", "Preserve observed existing-dev provenance; this is not a preregistered final assignment."])
        idx = splits_ws.max_row
        for cell in splits_ws[idx]: cell.alignment = wrap_top; cell.border = Border(bottom=thin_gray)
        splits_ws.cell(idx, 3).fill = PatternFill("solid", fgColor=yellow)
        splits_ws.cell(idx, 4).fill = PatternFill("solid", fgColor=yellow)
        splits_ws.row_dimensions[idx].height = 42
    for family in ("eacces-config", "wrong-model"):
        splits_ws.append([family, "existing-holdout", "excluded from current queue", None, "NO—known holdout", "pending", "Do not relabel as a new blind family."])
        idx = splits_ws.max_row
        for cell in splits_ws[idx]: cell.alignment = wrap_top; cell.border = Border(bottom=thin_gray)
    split_dv = DataValidation(type="list", formula1='"train,dev,blind"', allow_blank=True)
    splits_ws.add_data_validation(split_dv); split_dv.add(f"D2:D{splits_ws.max_row}")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    check = load_workbook(output_path, read_only=False, data_only=False)
    expected = {"Start Here", "Cases", "Fact Proposals", "Target Conflicts", "Family Split"}
    if set(check.sheetnames) != expected:
        raise RuntimeError(f"workbook sheets mismatch: {check.sheetnames}")
    if check["Cases"].max_row != len(queue) + 1 or check["Fact Proposals"].max_row != sum(len(r["facts"]) for r in proposals) + 1:
        raise RuntimeError("workbook row counts do not match source JSONL files")
    check.close()
    return output_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--queue", type=Path, default=DEFAULT_QUEUE)
    parser.add_argument("--proposals", type=Path, default=DEFAULT_PROPOSALS)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output = build_workbook(read_jsonl(args.queue), read_jsonl(args.proposals), args.out)
    print(f"[micro-generator-review-workbook] wrote {output.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
